from django.contrib import admin
from django.contrib.auth.admin import UserAdmin as BaseUserAdmin
from django.contrib.auth.models import User, Group

admin.site.unregister(Group)
from django.utils.html import format_html, mark_safe
from django.utils import timezone
from django.db.models import Avg, Count
from .models import Device, SensorData, Relay, RelayLog, DeviceConfig, ThresholdSetting, AIDecisionLog

# ─────────────────────────────────────────────
#  Admin Site Customization
# ─────────────────────────────────────────────
admin.site.site_header  = "🏠 IoT Dashboard Admin"
admin.site.site_title   = "IoT Dashboard"
admin.site.index_title  = "ระบบจัดการ IoT Dashboard"


# ─────────────────────────────────────────────
#  User Management (เพิ่ม / ลบ / แก้ไข Admin)
# ─────────────────────────────────────────────
admin.site.unregister(User)

@admin.register(User)
class CustomUserAdmin(BaseUserAdmin):
    list_display = ('username', 'email', 'first_name', 'last_name',
                    'is_active_badge', 'is_staff_badge', 'is_superuser_badge', 'date_joined')
    list_filter  = ('is_active', 'is_staff', 'is_superuser', 'date_joined')
    search_fields = ('username', 'email', 'first_name', 'last_name')
    ordering = ('-date_joined',)
    list_per_page = 20

    fieldsets = (
        ('ข้อมูลบัญชี', {
            'fields': ('username', 'password')
        }),
        ('ข้อมูลส่วนตัว', {
            'fields': ('first_name', 'last_name', 'email')
        }),
        ('สิทธิ์การเข้าถึง', {
            'fields': ('is_active', 'is_staff', 'is_superuser', 'groups', 'user_permissions'),
            'description': '⚠️ is_staff = เข้า Admin ได้  |  is_superuser = สิทธิ์สูงสุด'
        }),
        ('ประวัติ', {
            'fields': ('last_login', 'date_joined'),
            'classes': ('collapse',)
        }),
    )

    add_fieldsets = (
        ('สร้างผู้ใช้ใหม่', {
            'classes': ('wide',),
            'fields': ('username', 'email', 'first_name', 'last_name',
                       'password1', 'password2', 'is_active', 'is_staff', 'is_superuser'),
            'description': 'กรอกข้อมูลผู้ใช้ใหม่และกำหนดสิทธิ์การเข้าถึง'
        }),
    )

    actions = ['make_staff', 'remove_staff', 'activate_users', 'deactivate_users']

    def is_active_badge(self, obj):
        if obj.is_active:
            return format_html('<span style="color:green;font-weight:bold;">✅ Active</span>')
        return format_html('<span style="color:red;font-weight:bold;">❌ Inactive</span>')
    is_active_badge.short_description = 'สถานะ'

    def is_staff_badge(self, obj):
        if obj.is_staff:
            return format_html('<span style="color:#007bff;font-weight:bold;">🔑 Staff</span>')
        return format_html('<span style="color:#999;">—</span>')
    is_staff_badge.short_description = 'Staff'

    def is_superuser_badge(self, obj):
        if obj.is_superuser:
            return format_html('<span style="color:#e83e8c;font-weight:bold;">👑 Superuser</span>')
        return format_html('<span style="color:#999;">—</span>')
    is_superuser_badge.short_description = 'Superuser'

    @admin.action(description='✅ เปิดใช้งานบัญชี (Activate)')
    def activate_users(self, request, queryset):
        updated = queryset.update(is_active=True)
        self.message_user(request, f"เปิดใช้งาน {updated} บัญชีสำเร็จ")

    @admin.action(description='❌ ระงับบัญชี (Deactivate)')
    def deactivate_users(self, request, queryset):
        queryset = queryset.exclude(pk=request.user.pk)  # ป้องกันระงับตัวเอง
        updated = queryset.update(is_active=False)
        self.message_user(request, f"ระงับ {updated} บัญชีสำเร็จ")

    @admin.action(description='🔑 ให้สิทธิ์ Staff (เข้า Admin ได้)')
    def make_staff(self, request, queryset):
        updated = queryset.update(is_staff=True)
        self.message_user(request, f"ให้สิทธิ์ Staff {updated} บัญชีสำเร็จ")

    @admin.action(description='🔒 ถอนสิทธิ์ Staff')
    def remove_staff(self, request, queryset):
        queryset = queryset.exclude(pk=request.user.pk)  # ป้องกันถอนตัวเอง
        updated = queryset.update(is_staff=False)
        self.message_user(request, f"ถอนสิทธิ์ Staff {updated} บัญชีสำเร็จ")


# ─────────────────────────────────────────────
#  Device
# ─────────────────────────────────────────────
@admin.register(Device)
class DeviceAdmin(admin.ModelAdmin):
    list_display  = ('name', 'status_badge', 'last_updated_th', 'created_at_th')
    list_display_links = ('name',)
    list_filter   = ('is_on',)
    search_fields = ('name',)
    readonly_fields = ('created_at', 'last_updated')
    list_per_page = 20
    actions = ['turn_on_devices', 'turn_off_devices']

    class Media:
        css = {'all': ('iot_dashboard/admin_custom.css',)}

    def status_badge(self, obj):
        if obj.is_on:
            return mark_safe(
                '<span style="display:inline-block;padding:2px 10px;border-radius:12px;'
                'background:#d4edda;color:#155724;font-weight:bold;white-space:nowrap;">🟢 ON</span>'
            )
        return mark_safe(
            '<span style="display:inline-block;padding:2px 10px;border-radius:12px;'
            'background:#e2e3e5;color:#495057;white-space:nowrap;">⚫ OFF</span>'
        )
    status_badge.short_description = 'สถานะ'

    def last_updated_th(self, obj):
        return format_html(
            '<span style="white-space:nowrap;font-size:12px;color:#495057;">{}</span>',
            timezone.localtime(obj.last_updated).strftime('%d/%m/%Y %H:%M')
        )
    last_updated_th.short_description = 'อัปเดตล่าสุด'
    last_updated_th.admin_order_field = 'last_updated'

    def created_at_th(self, obj):
        return format_html(
            '<span style="white-space:nowrap;font-size:12px;color:#6c757d;">{}</span>',
            timezone.localtime(obj.created_at).strftime('%d/%m/%Y %H:%M')
        )
    created_at_th.short_description = 'สร้างเมื่อ'
    created_at_th.admin_order_field = 'created_at'

    @admin.action(description='🟢 เปิดอุปกรณ์ที่เลือก (Turn ON)')
    def turn_on_devices(self, request, queryset):
        updated = queryset.update(is_on=True, last_updated=timezone.now())
        self.message_user(request, f"เปิด {updated} อุปกรณ์สำเร็จ")

    @admin.action(description='⚫ ปิดอุปกรณ์ที่เลือก (Turn OFF)')
    def turn_off_devices(self, request, queryset):
        updated = queryset.update(is_on=False, last_updated=timezone.now())
        self.message_user(request, f"ปิด {updated} อุปกรณ์สำเร็จ")


# ─────────────────────────────────────────────
#  Relay
# ─────────────────────────────────────────────
@admin.register(Relay)
class RelayAdmin(admin.ModelAdmin):
    list_display  = ('name', 'relay_summary', 'relay1_badge', 'relay2_badge',
                     'relay3_badge', 'last_updated_th', 'created_at_th')
    list_display_links = ('name',)
    list_filter   = ('relay1_status', 'relay2_status', 'relay3_status')
    search_fields = ('name',)
    readonly_fields = ('created_at', 'last_updated', 'relay_panel')
    list_per_page = 20

    class Media:
        css = {'all': ('iot_dashboard/admin_custom.css',)}
    fieldsets = (
        ('ข้อมูล Relay', {
            'fields': ('name',)
        }),
        ('สถานะ Relay', {
            'fields': ('relay_panel', 'relay1_status', 'relay2_status', 'relay3_status'),
            'description': '✅ เปิด (True) = RELAY วงจรต่อ  |  ❌ ปิด (False) = RELAY วงจรตัด'
        }),
        ('Timestamps', {
            'fields': ('last_updated', 'created_at'),
            'classes': ('collapse',)
        }),
    )

    def _relay_pill(self, status, label):
        if status:
            return format_html(
                '<span style="display:inline-block;padding:3px 10px;border-radius:12px;'
                'background:#d4edda;color:#155724;font-weight:bold;font-size:12px;'
                'white-space:nowrap;">🟢 {}</span>',
                label
            )
        return format_html(
            '<span style="display:inline-block;padding:3px 10px;border-radius:12px;'
            'background:#e2e3e5;color:#495057;font-size:12px;white-space:nowrap;">'
            '⚫ {}</span>',
            label
        )

    def relay1_badge(self, obj):
        return self._relay_pill(obj.relay1_status, 'RELAY 1')
    relay1_badge.short_description = 'RELAY 1'

    def relay2_badge(self, obj):
        return self._relay_pill(obj.relay2_status, 'RELAY 2')
    relay2_badge.short_description = 'RELAY 2'

    def relay3_badge(self, obj):
        return self._relay_pill(obj.relay3_status, 'RELAY 3')
    relay3_badge.short_description = 'RELAY 3'

    def relay_summary(self, obj):
        on_count = sum([obj.relay1_status, obj.relay2_status, obj.relay3_status])
        if on_count == 0:
            color, text = '#6c757d', 'ทั้งหมดปิด'
        elif on_count == 3:
            color, text = '#155724', 'ทั้งหมดเปิด'
        else:
            color, text = '#856404', f'{on_count}/3 เปิด'
        return format_html(
            '<span style="font-weight:bold;color:{};">{}</span>',
            color, text
        )
    relay_summary.short_description = 'สรุป'

    def relay_panel(self, obj):
        """แสดง visual panel ของ relay ทั้ง 3 ตัว (read-only ใน detail view)"""
        relays = [
            ('RELAY 1', obj.relay1_status),
            ('RELAY 2', obj.relay2_status),
            ('RELAY 3', obj.relay3_status),
        ]
        cells = mark_safe(''.join(
            '<div style="display:inline-block;margin:4px 8px;text-align:center;">'
            f'<div style="width:70px;padding:8px 4px;border-radius:8px;font-weight:bold;font-size:13px;'
            f'background:{"#d4edda" if on else "#e2e3e5"};'
            f'color:{"#155724" if on else "#495057"};">'
            f'{"🟢" if on else "⚫"}</div>'
            f'<div style="font-size:11px;margin-top:4px;color:#6c757d;">{label}</div>'
            '</div>'
            for label, on in relays
        ))
        return format_html('<div style="display:flex;gap:4px;padding:4px 0;">{}</div>', cells)
    relay_panel.short_description = 'สถานะ Relay (ภาพรวม)'

    def last_updated_th(self, obj):
        return format_html(
            '<span style="white-space:nowrap;font-size:12px;color:#495057;">{}</span>',
            timezone.localtime(obj.last_updated).strftime('%d/%m/%Y %H:%M')
        )
    last_updated_th.short_description = 'อัปเดตล่าสุด'
    last_updated_th.admin_order_field = 'last_updated'

    def created_at_th(self, obj):
        return format_html(
            '<span style="white-space:nowrap;font-size:12px;color:#6c757d;">{}</span>',
            timezone.localtime(obj.created_at).strftime('%d/%m/%Y %H:%M')
        )
    created_at_th.short_description = 'สร้างเมื่อ'
    created_at_th.admin_order_field = 'created_at'

    def save_model(self, request, obj, form, change):
        """บันทึก RelayLog เมื่อมีการแก้ไขสถานะ relay ผ่าน Admin panel"""
        if change:
            try:
                old = Relay.objects.get(pk=obj.pk)
                relay_fields = [
                    (1, 'relay1_status'),
                    (2, 'relay2_status'),
                    (3, 'relay3_status'),
                ]
                user_info = request.user.username if request.user.is_authenticated else 'admin'
                for relay_num, field in relay_fields:
                    old_val = getattr(old, field)
                    new_val = getattr(obj, field)
                    if old_val != new_val:
                        RelayLog.record(
                            relay_number=relay_num,
                            new_state=new_val,
                            previous_state=old_val,
                            source='manual',
                            reason=f'Admin panel edited by {user_info}'
                        )
            except Relay.DoesNotExist:
                pass
        super().save_model(request, obj, form, change)


# ─────────────────────────────────────────────
#  Sensor Data
# ─────────────────────────────────────────────
@admin.register(SensorData)
class SensorDataAdmin(admin.ModelAdmin):
    list_display  = ('device_name_badge', 'temperature_badge', 'humidity_badge', 'ds18b20_badge', 'timestamp_th')
    list_display_links = ('device_name_badge',)
    list_filter   = ('device_name',)
    search_fields = ('device_name',)
    readonly_fields = ('timestamp',)
    date_hierarchy  = 'timestamp'
    ordering = ('-timestamp',)
    list_per_page = 50
    actions = ['delete_selected_data', 'delete_older_than_1day',
               'delete_older_than_7days', 'delete_older_than_30days', 'delete_all_data',
               'export_excel', 'export_json']

    class Media:
        css = {'all': ('iot_dashboard/admin_custom.css',)}

    def device_name_badge(self, obj):
        return format_html(
            '<span style="font-family:monospace;font-size:12px;color:#007bff;'
            'background:#e7f1ff;padding:2px 8px;border-radius:6px;white-space:nowrap;">{}</span>',
            obj.device_name
        )
    device_name_badge.short_description = 'Device'
    device_name_badge.admin_order_field = 'device_name'

    def timestamp_th(self, obj):
        return format_html(
            '<span style="white-space:nowrap;font-size:12px;color:#495057;">{}</span>',
            timezone.localtime(obj.timestamp).strftime('%d/%m/%Y %H:%M:%S')
        )
    timestamp_th.short_description = 'เวลา'
    timestamp_th.admin_order_field = 'timestamp'

    def temperature_badge(self, obj):
        if obj.temperature is None:
            return format_html('<span style="color:#999;">N/A</span>')
        color = '#dc3545' if obj.temperature > 30 else ('#007bff' if obj.temperature < 20 else '#28a745')
        icon  = '🔥' if obj.temperature > 30 else ('❄️' if obj.temperature < 20 else '🌡️')
        temp  = f'{obj.temperature:.1f}'
        return format_html(
            '<span style="color:{};font-weight:bold;">{} {}°C</span>',
            color, icon, temp
        )
    temperature_badge.short_description = 'อุณหภูมิ (XY-MD03)'

    def humidity_badge(self, obj):
        if obj.humidity is None:
            return format_html('<span style="color:#999;">N/A</span>')
        color = '#17a2b8' if obj.humidity > 70 else ('#fd7e14' if obj.humidity < 30 else '#28a745')
        icon  = '💧' if obj.humidity > 70 else ('🏜️' if obj.humidity < 30 else '💨')
        hum   = f'{obj.humidity:.1f}'
        return format_html(
            '<span style="color:{};font-weight:bold;">{} {}%</span>',
            color, icon, hum
        )
    humidity_badge.short_description = 'ความชื้น (XY-MD03)'

    def ds18b20_badge(self, obj):
        if obj.ds18b20_temperature is None:
            return format_html('<span style="color:#ccc;font-size:12px;">—</span>')
        temp = f'{obj.ds18b20_temperature:.1f}'
        color = '#dc3545' if obj.ds18b20_temperature > 35 else ('#007bff' if obj.ds18b20_temperature < 20 else '#6f42c1')
        return format_html(
            '<span style="color:{};font-weight:bold;white-space:nowrap;">🌡️ {}°C</span>',
            color, temp
        )
    ds18b20_badge.short_description = 'DS18B20'
    ds18b20_badge.admin_order_field = 'ds18b20_temperature'

    # ── Delete Actions ──────────────────────────────────────

    @admin.action(description='🗑️ ลบข้อมูลที่เลือก')
    def delete_selected_data(self, request, queryset):
        count = queryset.count()
        queryset.delete()
        self.message_user(request, f'✅ ลบข้อมูล {count} รายการสำเร็จ')

    @admin.action(description='📅 ลบข้อมูลเก่ากว่า 1 วัน')
    def delete_older_than_1day(self, request, queryset):
        from datetime import timedelta
        cutoff = timezone.now() - timedelta(days=1)
        count, _ = queryset.model.objects.filter(timestamp__lt=cutoff).delete()
        self.message_user(request, f'✅ ลบข้อมูลเก่ากว่า 1 วัน จำนวน {count} รายการสำเร็จ')

    @admin.action(description='📅 ลบข้อมูลเก่ากว่า 7 วัน')
    def delete_older_than_7days(self, request, queryset):
        from datetime import timedelta
        cutoff = timezone.now() - timedelta(days=7)
        count, _ = queryset.model.objects.filter(timestamp__lt=cutoff).delete()
        self.message_user(request, f'✅ ลบข้อมูลเก่ากว่า 7 วัน จำนวน {count} รายการสำเร็จ')

    @admin.action(description='📅 ลบข้อมูลเก่ากว่า 30 วัน')
    def delete_older_than_30days(self, request, queryset):
        from datetime import timedelta
        cutoff = timezone.now() - timedelta(days=30)
        count, _ = queryset.model.objects.filter(timestamp__lt=cutoff).delete()
        self.message_user(request, f'✅ ลบข้อมูลเก่ากว่า 30 วัน จำนวน {count} รายการสำเร็จ')

    @admin.action(description='⚠️ ลบข้อมูลทั้งหมด (ล้างตาราง)')
    def delete_all_data(self, request, queryset):
        count, _ = queryset.model.objects.all().delete()
        self.message_user(request, f'✅ ลบข้อมูลทั้งหมด จำนวน {count} รายการสำเร็จ')

    # ── Export Actions ──────────────────────────────────────

    @admin.action(description='📊 Export เป็น Excel (.xlsx)')
    def export_excel(self, request, queryset):
        import openpyxl
        from django.http import HttpResponse

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'SensorData'

        # Header row
        headers = ['ID', 'Device', 'อุณหภูมิ XY-MD03 (°C)', 'ความชื้น XY-MD03 (%)',
                   'DS18B20 (°C)', 'เวลา (Asia/Bangkok)']
        ws.append(headers)

        # Style header
        from openpyxl.styles import Font, PatternFill, Alignment
        header_fill = PatternFill('solid', fgColor='1F77B4')
        header_font = Font(bold=True, color='FFFFFF')
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        # Data rows
        for obj in queryset.order_by('-timestamp'):
            local_dt = timezone.localtime(obj.timestamp).strftime('%d/%m/%Y %H:%M:%S')
            ws.append([
                obj.pk,
                obj.device_name,
                obj.temperature,
                obj.humidity,
                obj.ds18b20_temperature,
                local_dt,
            ])

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 4

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="sensor_data.xlsx"'
        wb.save(response)
        return response

    @admin.action(description='📋 Export เป็น JSON (.json)')
    def export_json(self, request, queryset):
        import json
        from django.http import HttpResponse

        rows = []
        for obj in queryset.order_by('-timestamp'):
            rows.append({
                'id': obj.pk,
                'device_name': obj.device_name,
                'temperature_xymd03': obj.temperature,
                'humidity_xymd03': obj.humidity,
                'ds18b20_temperature': obj.ds18b20_temperature,
                'timestamp': timezone.localtime(obj.timestamp).strftime('%Y-%m-%d %H:%M:%S'),
            })

        content = json.dumps(rows, ensure_ascii=False, indent=2)
        response = HttpResponse(content, content_type='application/json; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="sensor_data.json"'
        return response


# ─────────────────────────────────────────────
#  Threshold Setting
# ─────────────────────────────────────────────
@admin.register(ThresholdSetting)
class ThresholdSettingAdmin(admin.ModelAdmin):
    list_display  = ('mode_badge', 'temperature_range', 'humidity_range',
                     'alarm_badge', 'relay1_auto_enabled', 'updated_at')
    readonly_fields = ('created_at', 'updated_at', 'last_triggered', 'alarm_reason', 'alarm_active')
    list_per_page = 10
    fieldsets = (
        ('โหมดควบคุม', {
            'fields': ('mode', 'relay1_auto_enabled'),
        }),
        ('ค่า Threshold อุณหภูมิ (°C)', {
            'fields': ('temperature_low', 'temperature_high'),
            'description': 'เมื่ออุณหภูมิออกนอกช่วงนี้จะกระตุ้น Alarm'
        }),
        ('ค่า Threshold ความชื้น (%)', {
            'fields': ('humidity_low', 'humidity_high'),
            'description': 'เมื่อความชื้นออกนอกช่วงนี้จะกระตุ้น Alarm'
        }),
        ('ตั้งค่า Hysteresis & Average', {
            'fields': ('hysteresis_percentage', 'average_window'),
            'classes': ('collapse',)
        }),
        ('สถานะ Alarm (อ่านอย่างเดียว)', {
            'fields': ('alarm_active', 'alarm_reason', 'last_triggered'),
            'classes': ('collapse',)
        }),
        ('Timestamps', {
            'fields': ('created_at', 'updated_at'),
            'classes': ('collapse',)
        }),
    )

    def mode_badge(self, obj):
        if obj.mode == 'AUTO':
            return format_html('<span style="color:#007bff;font-weight:bold;">🤖 AUTO</span>')
        return format_html('<span style="color:#6c757d;font-weight:bold;">👤 MANUAL</span>')
    mode_badge.short_description = 'โหมด'

    def temperature_range(self, obj):
        t_low  = f'{obj.temperature_low:.1f}'
        t_high = f'{obj.temperature_high:.1f}'
        return format_html(
            '<span style="color:#dc3545;">❄️ {}°C</span> → <span style="color:#dc3545;">🔥 {}°C</span>',
            t_low, t_high
        )
    temperature_range.short_description = 'ช่วงอุณหภูมิ'

    def humidity_range(self, obj):
        h_low  = f'{obj.humidity_low:.1f}'
        h_high = f'{obj.humidity_high:.1f}'
        return format_html(
            '<span style="color:#17a2b8;">🏜️ {}%</span> → <span style="color:#17a2b8;">💧 {}%</span>',
            h_low, h_high
        )
    humidity_range.short_description = 'ช่วงความชื้น'

    def alarm_badge(self, obj):
        if obj.alarm_active:
            return format_html('<span style="color:red;font-weight:bold;">🚨 ALARM</span>')
        return format_html('<span style="color:green;font-weight:bold;">✅ ปกติ</span>')
    alarm_badge.short_description = 'Alarm'


# ─────────────────────────────────────────────
#  Relay Log
# ─────────────────────────────────────────────
@admin.register(RelayLog)
class RelayLogAdmin(admin.ModelAdmin):
    list_display  = ('timestamp_th', 'relay_badge', 'state_badge', 'prev_state_badge',
                     'source_badge', 'reason_preview')
    list_display_links = ('timestamp_th',)
    list_filter   = ('relay_number', 'new_state', 'source', 'timestamp')
    search_fields = ('reason',)
    date_hierarchy  = 'timestamp'
    ordering = ('-timestamp',)
    list_per_page = 50
    actions = ['export_excel', 'export_json']

    class Media:
        css = {'all': ('iot_dashboard/admin_custom.css',)}

    def timestamp_th(self, obj):
        return format_html(
            '<span style="white-space:nowrap;font-size:12px;color:#495057;">{}</span>',
            timezone.localtime(obj.timestamp).strftime('%d/%m/%Y %H:%M:%S')
        )
    timestamp_th.short_description = 'เวลา'
    timestamp_th.admin_order_field = 'timestamp'

    def relay_badge(self, obj):
        colors = {1: '#007bff', 2: '#6f42c1', 3: '#fd7e14'}
        color = colors.get(obj.relay_number, '#6c757d')
        return format_html(
            '<span style="display:inline-block;padding:2px 10px;border-radius:12px;'
            'background:{0}20;color:{0};font-weight:bold;white-space:nowrap;'
            'font-size:12px;border:1px solid {0}40;">RELAY {1}</span>',
            color, obj.relay_number
        )
    relay_badge.short_description = 'Relay'
    relay_badge.admin_order_field = 'relay_number'

    def state_badge(self, obj):
        if obj.new_state:
            return format_html(
                '<span style="display:inline-block;padding:2px 10px;border-radius:12px;'
                'background:#d4edda;color:#155724;font-weight:bold;'
                'white-space:nowrap;">🟢 ON</span>'
            )
        return format_html(
            '<span style="display:inline-block;padding:2px 10px;border-radius:12px;'
            'background:#e2e3e5;color:#495057;white-space:nowrap;">⚫ OFF</span>'
        )
    state_badge.short_description = 'สถานะใหม่'
    state_badge.admin_order_field = 'new_state'

    def prev_state_badge(self, obj):
        if obj.previous_state is None:
            return format_html('<span style="color:#ccc;">-</span>')
        if obj.previous_state:
            return format_html('<span style="color:#155724;font-size:12px;">🟢 ON</span>')
        return format_html('<span style="color:#6c757d;font-size:12px;">⚫ OFF</span>')
    prev_state_badge.short_description = 'สถานะเดิม'
    prev_state_badge.admin_order_field = 'previous_state'

    SOURCE_STYLES = {
        'mqtt':      ('#0d6efd', '📡 ESP32'),
        'ai_agent':  ('#6f42c1', '🤖 AI Agent'),
        'threshold': ('#dc3545', '⚠️ Threshold'),
        'manual':    ('#198754', '✏️ Manual'),
    }

    def source_badge(self, obj):
        color, label = self.SOURCE_STYLES.get(obj.source, ('#6c757d', obj.source))
        return format_html(
            '<span style="display:inline-block;padding:2px 10px;border-radius:12px;'
            'background:{0}18;color:{0};font-weight:bold;white-space:nowrap;'
            'font-size:12px;border:1px solid {0}30;">{1}</span>',
            color, label
        )
    source_badge.short_description = 'แหล่งที่สั่ง'
    source_badge.admin_order_field = 'source'

    def reason_preview(self, obj):
        if not obj.reason:
            return format_html('<span style="color:#ccc;">-</span>')
        short = obj.reason[:80] + '…' if len(obj.reason) > 80 else obj.reason
        return format_html(
            '<span style="font-size:12px;color:#495057;" title="{full}">{short}</span>',
            full=obj.reason, short=short
        )
    reason_preview.short_description = 'เหตุผล'

    # ── Export Actions ───────────────────────────────────────

    @admin.action(description='📊 Export เป็น Excel (.xlsx)')
    def export_excel(self, request, queryset):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from django.http import HttpResponse

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'RelayLog'
        headers = ['ID', 'Relay', 'สถานะใหม่', 'สถานะเดิม', 'แหล่งที่สั่ง', 'เหตุผล', 'เวลา']
        ws.append(headers)
        header_fill = PatternFill('solid', fgColor='1F77B4')
        header_font = Font(bold=True, color='FFFFFF')
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        source_labels = dict(RelayLog.SOURCE_CHOICES)
        for obj in queryset.order_by('-timestamp'):
            ws.append([
                obj.pk,
                f'RELAY {obj.relay_number}',
                'ON' if obj.new_state else 'OFF',
                ('ON' if obj.previous_state else 'OFF') if obj.previous_state is not None else '-',
                source_labels.get(obj.source, obj.source),
                obj.reason,
                timezone.localtime(obj.timestamp).strftime('%d/%m/%Y %H:%M:%S'),
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="relay_log.xlsx"'
        wb.save(response)
        return response

    @admin.action(description='📋 Export เป็น JSON (.json)')
    def export_json(self, request, queryset):
        import json
        from django.http import HttpResponse

        source_labels = dict(RelayLog.SOURCE_CHOICES)
        rows = []
        for obj in queryset.order_by('-timestamp'):
            rows.append({
                'id': obj.pk,
                'relay_number': obj.relay_number,
                'new_state': 'ON' if obj.new_state else 'OFF',
                'previous_state': ('ON' if obj.previous_state else 'OFF') if obj.previous_state is not None else None,
                'source': source_labels.get(obj.source, obj.source),
                'reason': obj.reason,
                'timestamp': timezone.localtime(obj.timestamp).strftime('%Y-%m-%d %H:%M:%S'),
            })

        content = json.dumps(rows, ensure_ascii=False, indent=2)
        response = HttpResponse(content, content_type='application/json; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="relay_log.json"'
        return response


# ─────────────────────────────────────────────
#  AI Decision Log
# ─────────────────────────────────────────────
@admin.register(AIDecisionLog)
class AIDecisionLogAdmin(admin.ModelAdmin):
    list_display  = ('timestamp_th', 'decision_badge', 'confidence_bar',
                     'relay_status_badge', 'command_sent_badge', 'reasoning_preview')
    list_display_links = ('timestamp_th',)
    list_filter   = ('decision', 'relay_status', 'command_sent', 'timestamp')

    class Media:
        css = {'all': ('iot_dashboard/admin_custom.css',)}
    search_fields = ('reasoning',)
    readonly_fields = ('timestamp', 'decision', 'confidence', 'reasoning_box',
                       'weather_data_formatted', 'relay_status', 'command_sent',
                       'confidence_bar')
    date_hierarchy = 'timestamp'
    ordering = ('-timestamp',)
    list_per_page = 25
    show_full_result_count = False
    actions = ['delete_selected_logs']

    fieldsets = (
        ('🤖 ผลการตัดสินใจ AI', {
            'fields': ('timestamp', 'decision', 'confidence_bar', 'relay_status', 'command_sent'),
        }),
        ('📝 เหตุผลการตัดสินใจ', {
            'fields': ('reasoning_box',),
        }),
        ('🌤️ ข้อมูลสภาพอากาศที่ใช้', {
            'fields': ('weather_data_formatted',),
            'classes': ('collapse',),
        }),
    )

    # ── List view columns ──────────────────────────────────

    def timestamp_th(self, obj):
        return format_html(
            '<span style="white-space:nowrap;font-size:12px;">{}</span>',
            timezone.localtime(obj.timestamp).strftime('%d/%m/%Y %H:%M')
        )
    timestamp_th.short_description = 'เวลา'
    timestamp_th.admin_order_field = 'timestamp'

    def decision_badge(self, obj):
        if obj.decision == 'on':
            return format_html(
                '<span style="display:inline-block;padding:2px 8px;border-radius:12px;'
                'background:#d4edda;color:#155724;font-weight:bold;white-space:nowrap;">'
                '🟢 ON</span>'
            )
        return format_html(
            '<span style="display:inline-block;padding:2px 8px;border-radius:12px;'
            'background:#e2e3e5;color:#383d41;font-weight:bold;white-space:nowrap;">'
            '⚫ OFF</span>'
        )
    decision_badge.short_description = 'ตัดสินใจ'

    def confidence_bar(self, obj):
        pct   = int(obj.confidence * 100)
        pct_w = min(pct, 100)
        color = '#28a745' if pct >= 70 else ('#ffc107' if pct >= 40 else '#dc3545')
        pct_str = str(pct)
        pct_w_str = str(pct_w)
        return format_html(
            '<div style="min-width:90px;background:#e9ecef;border-radius:6px;overflow:hidden;">'
            '<div style="width:{}%;background:{};padding:2px 4px;text-align:center;'
            'color:white;font-size:11px;font-weight:bold;border-radius:6px;'
            'min-width:32px;white-space:nowrap;">{}%</div></div>',
            pct_w_str, color, pct_str
        )
    confidence_bar.short_description = 'ความมั่นใจ'

    def relay_status_badge(self, obj):
        if obj.relay_status:
            return format_html(
                '<span style="color:#155724;background:#d4edda;padding:2px 6px;'
                'border-radius:10px;font-size:12px;white-space:nowrap;">🟢 ON</span>'
            )
        return format_html(
            '<span style="color:#383d41;background:#e2e3e5;padding:2px 6px;'
            'border-radius:10px;font-size:12px;white-space:nowrap;">⚫ OFF</span>'
        )
    relay_status_badge.short_description = 'Relay 2'

    def command_sent_badge(self, obj):
        if obj.command_sent:
            return format_html(
                '<span style="color:#155724;font-size:14px;" title="ส่งคำสั่งสำเร็จ">✅</span>'
            )
        return format_html(
            '<span style="color:#721c24;font-size:14px;" title="ส่งคำสั่งไม่สำเร็จ">❌</span>'
        )
    command_sent_badge.short_description = 'ส่งสำเร็จ'

    def reasoning_preview(self, obj):
        text  = obj.reasoning[:80] + '…' if len(obj.reasoning) > 80 else obj.reasoning
        return format_html(
            '<span title="{}" style="display:block;max-width:300px;overflow:hidden;'
            'white-space:nowrap;text-overflow:ellipsis;font-size:12px;color:#495057;">{}</span>',
            obj.reasoning, text
        )
    reasoning_preview.short_description = 'เหตุผล AI'

    # ── Detail view fields ─────────────────────────────────

    def reasoning_box(self, obj):
        return format_html(
            '<div style="background:#f8f9fa;border:1px solid #dee2e6;border-radius:6px;'
            'padding:12px 16px;font-size:13px;line-height:1.7;color:#212529;'
            'max-width:700px;white-space:pre-wrap;">{}</div>',
            obj.reasoning
        )
    reasoning_box.short_description = 'เหตุผลการตัดสินใจ (เต็ม)'

    def weather_data_formatted(self, obj):
        import json
        data = obj.weather_data
        if not data:
            return format_html('<span style="color:#999;">ไม่มีข้อมูลสภาพอากาศ</span>')
        rows = mark_safe(''.join(
            '<tr>'
            f'<td style="padding:4px 12px;color:#6c757d;white-space:nowrap;font-size:12px;">{k}</td>'
            f'<td style="padding:4px 12px;font-weight:bold;font-size:12px;">{v}</td>'
            '</tr>'
            for k, v in (data.items() if isinstance(data, dict) else {})
        ))
        if not rows:
            pretty = json.dumps(data, ensure_ascii=False, indent=2)
            return format_html(
                '<pre style="background:#f8f9fa;border:1px solid #dee2e6;border-radius:6px;'
                'padding:12px;font-size:12px;max-width:600px;overflow:auto;">{}</pre>',
                pretty
            )
        return format_html(
            '<table style="border-collapse:collapse;background:#f8f9fa;'
            'border:1px solid #dee2e6;border-radius:6px;font-size:12px;">'
            '<thead><tr style="background:#e9ecef;">'
            '<th style="padding:6px 12px;text-align:left;">ข้อมูล</th>'
            '<th style="padding:6px 12px;text-align:left;">ค่า</th>'
            '</tr></thead><tbody>{}</tbody></table>',
            rows
        )
    weather_data_formatted.short_description = 'ข้อมูลสภาพอากาศ'

    @admin.action(description='🗑️ ลบ Log AI ที่เลือก')
    def delete_selected_logs(self, request, queryset):
        count = queryset.count()
        queryset.delete()
        self.message_user(request, f"ลบ AI Decision Log {count} รายการสำเร็จ")


# ─────────────────────────────────────────────
#  Device Config
# ─────────────────────────────────────────────
@admin.register(DeviceConfig)
class DeviceConfigAdmin(admin.ModelAdmin):
    list_display  = ('device_name_badge', 'mqtt_broker_badge', 'mqtt_port_badge',
                     'mqtt_client_id_preview', 'ds18b20_temp_display', 'updated_at_th')
    readonly_fields = ('updated_at', 'ds18b20_updated_at', 'ds18b20_temperature',
                       'all_topics_display')
    list_per_page = 10
    fieldsets = (
        ('Device Identity', {
            'fields': ('device_name',),
            'description': 'ต้องตรงกับ #define DEVICE_NAME ใน firmware'
        }),
        ('MQTT Broker', {
            'fields': ('mqtt_broker', 'mqtt_port', 'mqtt_client_id_prefix'),
        }),
        ('DS18B20 Sensor (อ่านอย่างเดียว)', {
            'fields': ('ds18b20_temperature', 'ds18b20_updated_at'),
            'classes': ('collapse',)
        }),
        ('MQTT Topics ทั้งหมด', {
            'fields': ('all_topics_display',),
            'classes': ('collapse',),
            'description': 'Topics ที่สร้างจาก Device Name ปัจจุบัน'
        }),
        ('Timestamps', {
            'fields': ('updated_at',),
            'classes': ('collapse',)
        }),
    )

    def device_name_badge(self, obj):
        return format_html(
            '<span style="font-family:monospace;font-weight:bold;color:#155724;'
            'background:#d4edda;padding:2px 8px;border-radius:6px;white-space:nowrap;">{}</span>',
            obj.device_name
        )
    device_name_badge.short_description = 'Device Name'
    device_name_badge.admin_order_field = 'device_name'

    def mqtt_broker_badge(self, obj):
        return format_html(
            '<span style="font-family:monospace;font-size:12px;white-space:nowrap;'
            'color:#495057;">{}</span>',
            obj.mqtt_broker
        )
    mqtt_broker_badge.short_description = 'MQTT Broker'
    mqtt_broker_badge.admin_order_field = 'mqtt_broker'

    def mqtt_port_badge(self, obj):
        return format_html(
            '<span style="font-family:monospace;font-size:12px;color:#007bff;'
            'font-weight:bold;">{}</span>',
            obj.mqtt_port
        )
    mqtt_port_badge.short_description = 'Port'
    mqtt_port_badge.admin_order_field = 'mqtt_port'

    def updated_at_th(self, obj):
        return format_html(
            '<span style="white-space:nowrap;font-size:12px;color:#495057;">{}</span>',
            timezone.localtime(obj.updated_at).strftime('%d/%m/%Y %H:%M')
        )
    updated_at_th.short_description = 'อัปเดตล่าสุด'
    updated_at_th.admin_order_field = 'updated_at'

    def mqtt_client_id_preview(self, obj):
        return format_html(
            '<span style="font-family:monospace;font-size:12px;color:#6c757d;'
            'white-space:nowrap;">{}</span>',
            f"{obj.mqtt_client_id_prefix}_{obj.device_name}"
        )
    mqtt_client_id_preview.short_description = 'MQTT Client ID'

    def ds18b20_temp_display(self, obj):
        if obj.ds18b20_temperature is not None:
            temp = f'{obj.ds18b20_temperature:.1f}'
            return format_html(
                '<span style="color:#dc3545;font-weight:bold;">🌡️ {}°C</span>',
                temp
            )
        return format_html('<span style="color:#999;">ไม่มีข้อมูล</span>')
    ds18b20_temp_display.short_description = 'DS18B20'

    def all_topics_display(self, obj):
        topics = obj.get_all_topics_display()
        rows = mark_safe(''.join(
            '<tr>'
            f'<td style="padding:4px 8px;color:#666;white-space:nowrap;">{k}</td>'
            f'<td style="padding:4px 8px;font-family:monospace;color:#007bff;">{v}</td>'
            '</tr>'
            for k, v in topics.items()
        ))
        return format_html(
            '<table style="border-collapse:collapse;font-size:12px;">{}</table>', rows
        )
    all_topics_display.short_description = 'MQTT Topics'
